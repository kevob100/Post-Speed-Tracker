"""Build the staffing schedules (data/<sport>/staffing/YYYY-MM-DD.yaml) from Google Sheets.

Two sheets, both readable by link, downloaded as .xlsx on every run:

  x_watcher        "NFL X Watcher": one tab per day ("10/4"), columns Tweets (L1-2),
                   Distributor (L3-5), Writers. The source from 2026-10-04 on.
  breaking_news    "RotoWire Breaking News Schedule": one tab per month, a block per day,
                   NFL and NFL News columns among the other sports. Fills every earlier day.

A cell lists names split by "/"; "Eric C. (10:30)/Adam" means Eric until 10:30, then Adam
(the {name, until} hand-off staffing.py reads). A time on any later name can't be read as
a hand-off, so that name counts for the whole hour. Names are cleaned against the sheet's
own Name Key plus the aliases in config ("Kevin" -> "KOB"), which also splits cells typed
without a slash ("CullumNick R.").

Generated files carry `source:` and are rewritten each run; a file without it is
hand-written and never touched. When both sheets cover a day, X Watcher wins.

Run: python -m src.staffing_import --sport nfl
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from pathlib import Path

import requests
import yaml

EXPORT = "https://docs.google.com/spreadsheets/d/{id}/export?format=xlsx"
DAYS = ("Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday")
X_ROLES = {"tweets": "Tweets (L1-2)", "distributor": "Distributor (L3-5)", "writers": "Writers"}
BN_ROLES = {"nfl": "NFL desk", "nfl_news": "NFL News"}
BN_COLUMNS = {"NFL": "nfl", "NFL News": "nfl_news"}


def slot_start(label) -> str | None:
    """'8-9 AM' -> '08:00', '12-1 PM' -> '12:00', '11-Noon' -> '11:00', '11-Mid' -> '23:00'."""
    m = re.match(r"\s*(\d{1,2})\s*-\s*(\d{1,2}\s*(AM|PM)|Noon|Mid)", str(label or ""), re.I)
    if not m:
        return None
    h, end = int(m.group(1)), m.group(2).lower()
    if end == "mid" or (end.endswith("pm") and h != 12):
        h += 12
    return f"{h:02d}:00"


class Names:
    """Canonical schedule names: the Name Key, cleaned case, aliases, missing slashes."""

    def __init__(self, known: list[str], aliases: dict[str, str] | None = None):
        self.aliases = {k.lower().rstrip("."): v for k, v in (aliases or {}).items()}
        self.known = {n.lower().rstrip("."): n for n in known}
        # Longest first, so "Nick R." wins over "Nick".
        self.order = sorted(self.known, key=len, reverse=True)

    def one(self, raw: str) -> str:
        n = re.sub(r"\s+", " ", raw).strip()
        key = n.lower().rstrip(".")
        return self.aliases.get(key) or self.known.get(key, n)

    def split(self, raw: str) -> list[str]:
        """A name, or several typed together without a slash."""
        n = re.sub(r"\s+", " ", raw).strip()
        if not n or n == "-":
            return []
        if n.lower().rstrip(".") in self.known or n.lower().rstrip(".") in self.aliases:
            return [self.one(n)]
        out, rest = [], n
        while rest:
            # A known name counts only where it ends: at the end, a period, or a new capital.
            hit = next((k for k in self.order if rest.lower().startswith(k)
                        and (len(rest) == len(k) or rest[len(k)] == "." or rest[len(k)].isupper())), None)
            if not hit:
                return out + [self.one(rest)]
            out.append(self.one(rest[:len(hit)]))
            rest = rest[len(hit):].lstrip(". ").strip()
        return out


def parse_cell(text, start: str, names: Names) -> list:
    """'Eric C. (10:30)/Adam' in the 10:00 slot -> [{name: Eric C., until: '10:30'}, Adam]."""
    out = []
    for i, part in enumerate(str(text or "").split("/")):
        m = re.search(r"\((\d{1,2}):(\d{2})\)", part)
        people = names.split(re.sub(r"\(.*?\)", "", part))
        if m and i == 0 and len(people) == 1 and "/" in str(text):
            out.append({"name": people[0], "until": f"{start[:2]}:{m.group(2)}"})
        else:
            out.extend(people)
    return out


def _date_for_tab(tab: str, today: date) -> date | None:
    """Tab '10/4' -> the date nearest today. The xlsx export drops the slash ('104'), so a
    digits-only name is tried as every month/day split ('111' is Jan 11 or Nov 1; the one
    nearer today wins)."""
    t = tab.strip()
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})", t)
    if m:
        splits = [(int(m.group(1)), int(m.group(2)))]
    elif re.fullmatch(r"\d{2,4}", t):
        splits = [(int(t[:i]), int(t[i:])) for i in (1, 2) if 0 < len(t) - i <= 2]
    else:
        return None
    cands = []
    for mo, d in splits:
        for y in (today.year - 1, today.year, today.year + 1):
            try:
                cands.append(date(y, mo, d))
            except ValueError:
                pass
    return min(cands, key=lambda c: abs((c - today).days)) if cands else None


def x_watcher(wb, names: Names, today: date) -> dict[str, dict]:
    out = {}
    for ws in wb.worksheets:
        day = _date_for_tab(ws.title, today)
        if not day:
            continue
        rows = list(ws.iter_rows(values_only=True))
        cols = {}
        for i, h in enumerate(rows[0] if rows else []):
            h = str(h or "").lower()
            role = "tweets" if "tweet" in h else "distributor" if "distrib" in h else \
                "writers" if "writer" in h else None
            if role:
                cols[role] = i
        slots = []
        for row in rows[1:]:
            start = slot_start(row[0] if row else None)
            if not start:
                continue
            slot = {"start": start}
            for role, i in cols.items():
                v = parse_cell(row[i] if i < len(row) else None, start, names)
                if v:
                    slot[role] = v
            if len(slot) > 1:
                slots.append(slot)
        if slots:
            out[day.isoformat()] = {"date": day.isoformat(), "roles": {r: X_ROLES[r] for r in cols},
                                    "slots": slots, "source": "x_watcher"}
    return out


def breaking_news(wb, names: Names, since: date) -> dict[str, dict]:
    out = {}
    for ws in wb.worksheets:
        if not re.fullmatch(r"[A-Z][a-z]+ \d{2}", ws.title):   # month tabs; skip "Copy of ..."
            continue
        day, cols, slots = None, {}, []

        def flush():
            if day and slots and day >= since:
                out[day.isoformat()] = {"date": day.isoformat(),
                                        "roles": {r: BN_ROLES[r] for r in cols.values()},
                                        "slots": list(slots), "source": "breaking_news"}
        for row in ws.iter_rows(values_only=True):
            first = row[0] if row else None
            if isinstance(first, datetime):
                flush()
                day, cols, slots = first.date(), {}, []
            elif first in DAYS:
                cols = {i: BN_COLUMNS[h] for i, h in enumerate(row) if h in BN_COLUMNS}
            elif day and cols and (start := slot_start(first)):
                slot = {"start": start}
                for i, role in cols.items():
                    v = parse_cell(row[i] if i < len(row) else None, start, names)
                    if v:
                        slot[role] = v
                if len(slot) > 1:
                    slots.append(slot)
        flush()
    return out


def name_key(wb) -> list[str]:
    if "Name Key" not in wb.sheetnames:
        return []
    return [str(r[0]).strip() for r in list(wb["Name Key"].iter_rows(values_only=True))[1:] if r and r[0]]


def _download(sheet_id: str):
    import openpyxl
    resp = requests.get(EXPORT.format(id=sheet_id), timeout=60)
    resp.raise_for_status()
    return openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True)


def write(schedules: dict[str, dict], folder: Path, timezone: str) -> dict:
    """Write each day's file; leaves hand-written files (no `source`) alone."""
    folder.mkdir(parents=True, exist_ok=True)
    written = kept = 0
    for d, sched in sorted(schedules.items()):
        path = folder / f"{d}.yaml"
        if path.exists() and "source" not in (yaml.safe_load(path.read_text()) or {}):
            kept += 1
            continue
        body = {"date": d, "timezone": timezone, "source": sched["source"],
                "roles": sched["roles"], "slots": sched["slots"]}
        text = (f"# Generated from the {sched['source']} staffing sheet; edit the sheet, not this file.\n"
                + yaml.safe_dump(body, sort_keys=False, default_flow_style=None, width=200))
        if not path.exists() or path.read_text() != text:
            path.write_text(text)
            written += 1
    return {"days": len(schedules), "written": written, "hand_written_kept": kept}


def sync(cfg: dict, folder: Path, today: date | None = None) -> dict:
    """Download the configured sheets and rewrite the generated schedule files."""
    today = today or date.today()
    aliases = cfg.get("aliases") or {}
    days: dict[str, dict] = {}
    bn_id = cfg.get("breaking_news_sheet")
    known: list[str] = []
    if bn_id:
        wb = _download(bn_id)
        known = name_key(wb)
        since = date.fromisoformat(str(cfg.get("since") or "2026-04-01"))
        days.update(breaking_news(wb, Names(known, aliases), since))
    if cfg.get("x_watcher_sheet"):
        wb = _download(cfg["x_watcher_sheet"])
        days.update(x_watcher(wb, Names(known, aliases), today))   # X Watcher wins
    # The month tabs are planned weeks ahead; a day with no news yet only clutters the tab.
    days = {d: v for d, v in days.items() if d <= today.isoformat()}
    return write(days, folder, cfg.get("timezone") or "America/New_York")


if __name__ == "__main__":
    import argparse

    from .config import load_config, sport_data_dir

    ap = argparse.ArgumentParser(description="Rebuild staffing schedules from the Google Sheets.")
    ap.add_argument("--sport", default="nfl")
    args = ap.parse_args()
    conf = load_config()["sports"][args.sport]["staffing_sheets"]
    print(sync(conf, sport_data_dir(args.sport) / "staffing"))
